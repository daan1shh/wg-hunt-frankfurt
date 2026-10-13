"""Read and write WG tracker.xlsx. Dedupe by Listing ID.

    python src/tracker.py init-template            # write tracker_template.xlsx (headers only)
    python src/tracker.py migrate                  # add/format columns in WG tracker.xlsx
    python src/tracker.py has <id>                 # exit 0 if the listing is already tracked
    python src/tracker.py upsert <id> [<id> ...]   # write rows from listings/<id>.json, scored
    python src/tracker.py set <id> Col=value ...   # e.g. Status=Sent "Sent on=2026-11-02"
    python src/tracker.py list [--status S]        # print rows for /wg-status
"""
import datetime as dt
import shutil
import sys

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList

from common import ROOT, load_config, load_listing, to_date
from score import score

TRACKER = ROOT / "WG tracker.xlsx"
TEMPLATE = ROOT / "tracker_template.xlsx"
LOCK = ROOT / "~$WG tracker.xlsx"
BACKUPS = ROOT / "backups"

# (header, width, kind). The first 8 are the original columns and stay first.
COLUMNS = [
    ("WG address", 30, "text"),
    ("Roommates", 10, "int"),
    ("Dist from DWS", 11, "km"),
    ("Availability", 24, "text"),
    ("Rent", 9, "eur"),
    ("Caution", 9, "eur"),
    ("Misc costs", 9, "eur"),
    ("Room size", 9, "m2"),
    ("Listing ID", 10, "text"),
    ("URL", 14, "url"),
    ("Type", 12, "text"),
    ("Ad title", 32, "text"),
    ("Stadtteil", 14, "text"),
    ("Date found", 11, "date"),
    ("Ad online since", 11, "date"),
    ("Kaltmiete", 9, "eur"),
    ("Nebenkosten", 9, "eur"),
    ("Sonstige Kosten", 9, "eur"),
    ("Gesamtmiete", 10, "eur"),
    ("Ablöse", 9, "eur"),
    ("Available from", 11, "date"),
    ("Available until", 11, "date"),
    ("Covers Dec–May?", 9, "text"),
    ("Zwischenmiete?", 9, "text"),
    ("Flatmates (count, ages, genders)", 24, "text"),
    ("WG type", 13, "text"),
    ("WG languages", 13, "text"),
    ("Smoking", 10, "text"),
    ("Furnished", 10, "text"),
    ("Anmeldung possible", 10, "text"),
    ("Transit time to DWS", 9, "min"),
    ("Bike time to DWS", 9, "min"),
    ("Walk time to DWS", 9, "min"),
    ("Nearest stations", 30, "text"),
    ("Amenities", 40, "text"),
    ("Contact name(s)", 16, "text"),
    ("Ad language", 8, "text"),
    ("Trick question", 30, "text"),
    ("Score", 7, "int"),
    ("Status", 16, "text"),
    ("Draft file", 18, "text"),
    ("Sent on", 11, "date"),
    ("Follow-up on", 11, "date"),
    ("Notes", 50, "text"),
]
HEADERS = [c[0] for c in COLUMNS]
NUMBER_FORMATS = {"eur": '#,##0 "€"', "km": '0.0 "km"', "m2": '0 "m²"',
                  "min": '0 "min"', "date": "DD.MM.YYYY", "int": "0"}

STATUSES = ["New", "Drafted", "On hold (trick Q)", "Needs review", "Approved", "Sent",
            "Replied", "Viewing", "Rejected", "Got it"]
STATUS_FILLS = {
    "New": "DDEBF7", "Drafted": "FFF2CC", "On hold (trick Q)": "F8CBAD",
    "Needs review": "FCE4D6", "Approved": "C6E0B4", "Sent": "BDD7EE",
    "Replied": "A9D08E", "Viewing": "70AD47", "Rejected": "D9D9D9", "Got it": "00B050",
}
# statuses the scanner may overwrite on re-upsert; anything later is a human decision
AUTO_STATUSES = {None, "", "New", "On hold (trick Q)", "Needs review", "Rejected"}
MANUAL_COLUMNS = {"Draft file", "Sent on", "Follow-up on"}


# ---------- workbook helpers ----------

def check_lock():
    if LOCK.exists():
        sys.exit("WG tracker.xlsx is open in Excel (~$WG tracker.xlsx exists). Close Excel and retry.")


def backup():
    BACKUPS.mkdir(exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(TRACKER, BACKUPS / f"WG tracker.{stamp}.xlsx")
    for old in sorted(BACKUPS.glob("WG tracker.*.xlsx"))[:-20]:
        old.unlink()


def header_map(ws):
    return {c.value: c.column for c in ws[1] if c.value}


def format_sheet(ws):
    """Bold frozen header, filter, widths, number formats, Status colour rules."""
    cols = header_map(ws)
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    last = get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A1:{last}{max(ws.max_row, 2)}"
    for name, width, kind in COLUMNS:
        col = cols.get(name)
        if not col:
            continue
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = width
        fmt = NUMBER_FORMATS.get(kind)
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if fmt:
                cell.number_format = fmt
            cell.alignment = Alignment(wrap_text=kind == "text" and width >= 24, vertical="top")
    status_col = cols.get("Status")
    if status_col:
        sl = get_column_letter(status_col)
        rng = f"{sl}2:{sl}1000"
        ws.data_validations = DataValidationList()
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        dv.add(f"{sl}2:{sl}1000")
        ws.add_data_validation(dv)
        ws.conditional_formatting = type(ws.conditional_formatting)()
        for status, colour in STATUS_FILLS.items():
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'${sl}2="{status}"'],
                fill=PatternFill(start_color=colour, end_color=colour, fill_type="solid")))


def open_tracker():
    check_lock()
    wb = openpyxl.load_workbook(TRACKER)
    ws = wb.worksheets[0]
    missing = [h for h in HEADERS if h not in header_map(ws)]
    if missing:
        sys.exit(f"Tracker is missing columns {missing[:3]}... Run: python src/tracker.py migrate")
    return wb, ws


def save(wb):
    check_lock()
    backup()
    format_sheet(wb.worksheets[0])
    wb.save(TRACKER)


def find_row(ws, listing_id):
    col = header_map(ws)["Listing ID"]
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=col).value or "") == str(listing_id):
            return row
    return None


# ---------- listing -> row ----------

def _fmt_date(d):
    d = to_date(d)
    return d.strftime("%d.%m.%Y") if d else None


def row_from_listing(L, result):
    avail = f"ab {_fmt_date(L.get('available_from')) or 'sofort'}"
    avail += f" bis {_fmt_date(L.get('available_until'))}" if L.get("available_until") else ", unbefristet"
    notes = result["why"]
    if L.get("approx_location"):
        notes += " | distance approx (no exact address)"
    if L.get("hold_reasons"):
        notes += " | HOLD: " + "; ".join(L["hold_reasons"])
    if L.get("notes"):
        notes += " | " + L["notes"]
    yn = lambda v: None if v is None else ("Yes" if v is True else "No" if v is False else v)
    return {
        "WG address": L.get("address"),
        "Roommates": L.get("flatmates_count"),
        "Dist from DWS": L.get("dist_km"),
        "Availability": avail,
        "Rent": L.get("gesamtmiete"),
        "Caution": L.get("kaution"),
        "Misc costs": L.get("sonstige_kosten"),
        "Room size": L.get("room_size"),
        "Listing ID": str(L["id"]),
        "URL": L.get("url"),
        "Type": L.get("type"),
        "Ad title": L.get("title"),
        "Stadtteil": L.get("stadtteil"),
        "Date found": to_date(L.get("date_found")) or dt.date.today(),
        "Ad online since": to_date(L.get("online_since")),
        "Kaltmiete": L.get("kaltmiete"),
        "Nebenkosten": L.get("nebenkosten"),
        "Sonstige Kosten": L.get("sonstige_kosten"),
        "Gesamtmiete": L.get("gesamtmiete"),
        "Ablöse": L.get("abloese"),
        "Available from": to_date(L.get("available_from")),
        "Available until": to_date(L.get("available_until")),
        "Covers Dec–May?": result["covers"],
        "Zwischenmiete?": yn(L.get("zwischenmiete")),
        "Flatmates (count, ages, genders)": L.get("flatmates_detail"),
        "WG type": L.get("wg_type"),
        "WG languages": L.get("wg_languages"),
        "Smoking": L.get("smoking"),
        "Furnished": yn(L.get("furnished")),
        "Anmeldung possible": L.get("anmeldung") or "Unknown",
        "Transit time to DWS": L.get("transit_min"),
        "Bike time to DWS": L.get("bike_min"),
        "Walk time to DWS": L.get("walk_min") if (L.get("walk_min") or 99) < 30 else None,
        "Nearest stations": L.get("stations"),
        "Amenities": L.get("amenities"),
        "Contact name(s)": ", ".join(L.get("contacts") or []) or None,
        "Ad language": L.get("ad_language"),
        "Trick question": L.get("trick_question"),
        "Score": result["score"],
        "Notes": notes,
    }


def auto_status(L, result):
    if result["rejected"]:
        return "Rejected"
    if L.get("trick_question") or L.get("hold_reasons"):
        return "On hold (trick Q)"
    if result["covers"] == "Partly" or (L.get("anmeldung") or "").lower() == "no":
        return "Needs review"
    return "New"


# ---------- commands ----------

def cmd_init_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tracker"
    ws.append(HEADERS)
    format_sheet(ws)
    wb.save(TEMPLATE)
    print(f"wrote {TEMPLATE.name} with {len(HEADERS)} columns")


def cmd_migrate():
    check_lock()
    wb = openpyxl.load_workbook(TRACKER)
    ws = wb.worksheets[0]
    have = header_map(ws)
    for i, h in enumerate(HEADERS[:8], start=1):
        if ws.cell(row=1, column=i).value != h:
            sys.exit(f"Column {i} is {ws.cell(row=1, column=i).value!r}, expected {h!r}. Not touching it.")
    added = []
    for h in HEADERS:
        if h not in have:
            ws.cell(row=1, column=ws.max_column + 1, value=h)
            added.append(h)
    save(wb)
    print(f"added {len(added)} columns, formatted sheet, backup in backups/")


def cmd_has(listing_id):
    check_lock()
    wb = openpyxl.load_workbook(TRACKER, read_only=True)
    ws = wb.worksheets[0]
    cols = {c.value: i for i, c in enumerate(next(ws.iter_rows(max_row=1)))}
    i = cols.get("Listing ID")
    found = i is not None and any(str(r[i] or "") == str(listing_id)
                                  for r in ws.iter_rows(min_row=2, values_only=True))
    print("tracked" if found else "new")
    sys.exit(0 if found else 1)


def cmd_upsert(ids):
    cfg = load_config()
    wb, ws = open_tracker()
    cols = header_map(ws)
    for listing_id in ids:
        L = load_listing(listing_id)
        result = score(L, cfg)
        values = row_from_listing(L, result)
        row = find_row(ws, L["id"])
        if row is None:
            row = ws.max_row + 1
            values["Status"] = auto_status(L, result)
            verb = "added"
        else:
            current = ws.cell(row=row, column=cols["Status"]).value
            if current in AUTO_STATUSES:
                values["Status"] = auto_status(L, result)
            verb = "updated"
        for name, value in values.items():
            if name in MANUAL_COLUMNS:
                continue
            cell = ws.cell(row=row, column=cols[name], value=value)
            if name == "URL" and value:
                cell.hyperlink = value
                cell.value = "open ad"
                cell.style = "Hyperlink"
        status = ws.cell(row=row, column=cols["Status"]).value
        print(f"{verb} {L['id']}: {status}, {result['why']}")
    save(wb)


def cmd_set(listing_id, pairs):
    cfg = load_config()
    wb, ws = open_tracker()
    cols = header_map(ws)
    row = find_row(ws, listing_id)
    if row is None:
        sys.exit(f"{listing_id} not in tracker")
    for pair in pairs:
        name, _, value = pair.partition("=")
        if name not in cols:
            sys.exit(f"unknown column {name!r}")
        if name == "Status" and value not in STATUSES:
            sys.exit(f"unknown status {value!r}. One of: {STATUSES}")
        kind = dict((c[0], c[2]) for c in COLUMNS).get(name)
        if kind == "date" and value:
            value = to_date(value)
        elif kind in ("eur", "int", "min", "km", "m2") and value:
            value = float(value)
        ws.cell(row=row, column=cols[name], value=value or None)
        if name == "Sent on" and value:
            follow = value + dt.timedelta(days=cfg["sending"]["follow_up_days"])
            ws.cell(row=row, column=cols["Follow-up on"], value=follow)
    save(wb)
    print(f"{listing_id}: set {', '.join(pairs)}")


def cmd_list(status=None):
    check_lock()
    wb = openpyxl.load_workbook(TRACKER, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return
    head = rows[0]
    idx = {h: i for i, h in enumerate(head)}
    show = ["Listing ID", "Score", "Status", "Gesamtmiete", "Roommates", "Transit time to DWS",
            "Stadtteil", "Type", "Covers Dec–May?", "Sent on", "Follow-up on", "Draft file", "Notes"]
    out = []
    for r in rows[1:]:
        if not any(r):
            continue
        rec = {h: r[idx[h]] if h in idx and idx[h] < len(r) else None for h in show}
        if status and rec["Status"] != status:
            continue
        out.append(rec)
    out.sort(key=lambda x: -(x["Score"] or 0))
    for rec in out:
        print(" | ".join("" if rec[h] is None else (rec[h].strftime("%d.%m") if isinstance(rec[h], dt.datetime) else str(rec[h])) for h in show))
    print(f"\n{len(out)} rows")


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        sys.exit(__doc__)
    cmd, rest = args[0], args[1:]
    if cmd == "init-template":
        cmd_init_template()
    elif cmd == "migrate":
        cmd_migrate()
    elif cmd == "has" and len(rest) == 1:
        cmd_has(rest[0])
    elif cmd == "upsert" and rest:
        cmd_upsert(rest)
    elif cmd == "set" and len(rest) >= 2:
        cmd_set(rest[0], rest[1:])
    elif cmd == "list":
        cmd_list(rest[1] if rest[:1] == ["--status"] and len(rest) > 1 else None)
    else:
        sys.exit(__doc__)
